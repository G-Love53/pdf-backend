#!/usr/bin/env python3
"""Parse Jon's GUARD UAT xlsm → JSON for scripts/run-guard-uat.mjs."""

import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

DEFAULT_XLSM = Path.home() / "Downloads/CID GUARD WC Q Changes - UAT Test Pack v1.xlsm"
OUT = Path(__file__).resolve().parent.parent / "data/guard-uat-cases-v1.json"

CLASS_SEGMENT = {
    "5183CO": {"segment": "plumber", "ratingClassificationCd": "518322"},
    "5474CO": {"segment": "plumber", "ratingClassificationCd": "547400"},
    "5537CO": {"segment": "hvac", "ratingClassificationCd": "553700"},
    "5190CO": {"segment": "electrical", "ratingClassificationCd": "519000"},
    "9014CO": {"segment": "cleaning", "ratingClassificationCd": "901400"},
    "0917CO": {"segment": "pet", "ratingClassificationCd": "091700"},
    "9063CO": {"segment": "fitness_trainer", "ratingClassificationCd": "906300"},
    "9586CO": {"segment": "beauty_hair", "ratingClassificationCd": "958600"},
}

ENTITY_MAP = {
    "corporation": "CP",
    "llc": "LL",
    "individual": "IN",
    "partnership": "PT",
}

SKIP_Q = {
    "Are ACORD questions all answered favorably?",
    "Class-Specific UW Questions",
    "Do UW question responses appropriately map?",
    "Does the experience mod map into the submission?",
}


def col(row, idx):
    return str(row[idx] or "").strip() if len(row) > idx else ""


def parse_address(raw):
    s = str(raw or "").strip()
    m = re.match(r"^(.+?),\s*([^,]+),\s*([A-Z]{2})\s+(\d{5})", s)
    if not m:
        return {"street": s, "city": "Denver", "state": "CO", "zip": "80202"}
    return {
        "street": m.group(1).strip(),
        "city": m.group(2).strip(),
        "state": m.group(3).strip(),
        "zip": m.group(4).strip(),
    }


def is_case_start(row):
    try:
        num = int(float(row[0]))
    except (TypeError, ValueError):
        return None
    outcome = col(row, 1)
    cls = col(row, 2)
    if outcome in ("Quote", "Refer", "Decline") and cls.endswith("CO"):
        return num, outcome, cls
    return None


def parse_sheet(ws, sheet_name):
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    starts = []
    for i, row in enumerate(rows):
        hit = is_case_start(row)
        if hit:
            starts.append((i, hit[0], hit[1], hit[2]))

    cases = []
    for idx, (start_i, case_num, test_type, class_code) in enumerate(starts):
        end_i = starts[idx + 1][0] if idx + 1 < len(starts) else len(rows)
        header = rows[start_i]
        block = rows[start_i + 1 : end_i]
        case = {
            "id": f"{sheet_name}-{case_num}",
            "sheet": sheet_name,
            "number": case_num,
            "expectedOutcome": test_type,
            "classCode": class_code,
            "questions": [],
            "notes": [],
            "acordFavorable": True,
        }

        hq = col(header, 4)
        hans = col(header, 5)
        hexp = col(header, 6)
        if hq == "Are ACORD questions all answered favorably?":
            case["acordFavorable"] = hans.lower().startswith("y")
            if hexp:
                case["expectedReason"] = hexp
        elif hq and hq not in SKIP_Q and hq != "**None":
            case["questions"].append(
                {"text": hq, "answer": hans, "expectedNote": hexp or None}
            )

        def record_question(q, ans, exp):
            if q == "Are ACORD questions all answered favorably?":
                case["acordFavorable"] = ans.lower().startswith("y")
                if exp:
                    case["expectedReason"] = exp
            elif q and q not in SKIP_Q and q != "**None":
                case["questions"].append(
                    {"text": q, "answer": ans, "expectedNote": exp or None}
                )
            if exp and ("Referral" in exp or "Declination" in exp):
                case.setdefault("expectedReason", exp)

        pending = None
        for r in block:
            b = col(r, 1)
            q = col(r, 4)
            ans = col(r, 5)
            exp = col(r, 6)
            record_question(q, ans, exp)

            if b == "Address":
                pending = "address"
                case.setdefault("addresses", [])
                continue
            if pending == "address":
                if b and "," in b:
                    addr = parse_address(b)
                    case["addresses"].append(addr)
                    case["address"] = case["addresses"][0]
                    continue
                pending = None

            if b == "Payroll":
                pending = "payroll"
                continue
            if pending == "payroll":
                try:
                    case["payroll"] = int(float(b))
                except (TypeError, ValueError):
                    pass
                pending = None
                continue
            if b == "Entity":
                pending = "entity"
                continue
            if pending == "entity" and b:
                case["entity"] = ENTITY_MAP.get(b.lower(), "LL")
                case["entityLabel"] = b
                pending = None
                continue
            if b.startswith("Ex Mod"):
                pending = "exmod"
                continue
            if pending == "exmod":
                try:
                    case["experienceMod"] = float(b)
                except (TypeError, ValueError):
                    pass
                pending = None
                continue
            if b.startswith("*"):
                case["notes"].append(b)

        meta = CLASS_SEGMENT.get(class_code, {})
        case["segment"] = meta.get("segment", "plumber")
        case["ratingClassificationCd"] = meta.get(
            "ratingClassificationCd", class_code.replace("CO", "") + "00"
        )
        case["ownerIncluded"] = any(
            "officer" in n.lower() or "include officer" in n.lower()
            for n in case["notes"]
        )
        cases.append(case)
    return cases


def main():
    xlsm = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSM
    if not xlsm.is_file():
        print(f"Missing {xlsm}", file=sys.stderr)
        sys.exit(1)
    wb = openpyxl.load_workbook(xlsm, read_only=True, data_only=True)
    all_cases = []
    for name in ("Contractors", "Non-Contractors"):
        if name in wb.sheetnames:
            all_cases.extend(parse_sheet(wb[name], name))
    wb.close()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"source": xlsm.name, "cases": all_cases}, indent=2))
    print(f"Wrote {len(all_cases)} cases → {OUT}")


if __name__ == "__main__":
    main()
